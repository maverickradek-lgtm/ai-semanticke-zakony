"""
Rozhodnicek: sber usneseni vlady CR (odok.gov.cz) a usneseni Poslanecke snemovny (psp.cz)
do SAMOSTATNE Neon databaze (secret NEON_ROZHODNICEK_DB_URL). Nesmi se michat s databazi zakonu.

Prostredi:
  NEON_ROZHODNICEK_DB_URL  pripojeni k databazi (povinne)
  ROZHODNICEK_YEARS        roky oddelene carkou, napr. "2025" nebo "2024,2025" (vychozi: aktualni rok)
  ROZHODNICEK_SOURCES      "vlada,ps" (vychozi obe)
  ROZHODNICEK_DRY          "1" = nic nezapisovat, jen vypsat co by se stalo
  TIME_BUDGET_SECONDS      casovy limit behu (vychozi 12000)

Embeddingy se tady nepocitaji (chunks.embedding zustava NULL).
"""
import hashlib
import io
import os
import re
import sys
import time
import datetime as dt

import requests
import psycopg2
import psycopg2.extras
from bs4 import BeautifulSoup

ODOK = "https://odok.gov.cz"
PSP = "https://www.psp.cz"
UA = {"User-Agent": "Mozilla/5.0 (compatible; Rozhodnicek/1.0; +https://paragralf.cz)"}

START = time.time()
BUDGET = int(os.environ.get("TIME_BUDGET_SECONDS", "12000"))
DRY = os.environ.get("ROZHODNICEK_DRY", "") == "1"

MONTHS = {
    "ledna": 1, "února": 2, "března": 3, "dubna": 4, "května": 5, "června": 6,
    "července": 7, "srpna": 8, "září": 9, "října": 10, "listopadu": 11, "prosince": 12,
}


def log(*a):
    print(*a, flush=True)


def time_left():
    return BUDGET - (time.time() - START)


SESSION = requests.Session()
SESSION.headers.update(UA)


def get(url, binary=False, retries=4):
    last = None
    for i in range(retries):
        try:
            r = SESSION.get(url, timeout=60)
            if r.status_code == 404:
                return None
            r.raise_for_status()
            time.sleep(0.25)
            return r.content if binary else r.text
        except Exception as e:  # noqa
            last = e
            time.sleep(2 * (i + 1))
    log("  ! stazeni selhalo: " + url + " (" + str(last) + ")")
    return None


def docx_text(data):
    from docx import Document
    d = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = []
            for c in row.cells:
                tx = c.text.strip()
                if tx and tx not in cells:
                    cells.append(tx)
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def pdf_text(data):
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    return "\n".join((pg.extract_text() or "").strip() for pg in r.pages).strip()


def clean(text):
    text = text.replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def make_chunks(title, text, size=1400):
    paras = [p.strip() for p in text.split("\n") if p.strip()]
    chunks, cur = [], ""
    for p in paras:
        if cur and len(cur) + len(p) + 1 > size:
            chunks.append(cur)
            cur = p
        else:
            cur = (cur + "\n" + p) if cur else p
    if cur:
        chunks.append(cur)
    if not chunks:
        chunks = [title or ""]
    # prvni chunk dostane nazev, aby ho bylo mozne najit podle tematu
    if title and title not in chunks[0]:
        chunks[0] = title + "\n" + chunks[0]
    return chunks


# ---------------------------------------------------------------- vlada
def vlada_resolutions(year):
    """vraci seznam (cislo, datum jednani) z odok.gov.cz pro dany rok"""
    html = get(ODOK + "/portal/zvlady/jednani/" + str(year) + "/")
    if not html:
        return []
    dates = sorted(set(re.findall(r"/jednani-detail/(\d{4}-\d{2}-\d{2})/", html)))
    out = []
    for d in dates:
        det = get(ODOK + "/portal/zvlady/jednani-detail/" + d + "/")
        if not det:
            continue
        nums = sorted(set(int(n) for n in re.findall(r"/zvlady/usneseni/" + str(year) + r"/(\d+)/", det)))
        log("  jednani " + d + ": " + str(len(nums)) + " usneseni")
        for n in nums:
            out.append((n, d))
    return out


def vlada_fetch(year, cislo, datum):
    page = get(ODOK + "/portal/zvlady/usneseni/" + str(year) + "/" + str(cislo) + "/")
    nazev, cj, pid = "", "", ""
    if page:
        soup = BeautifulSoup(page, "html.parser")
        main = soup.find("main") or soup
        txt = main.get_text("\n", strip=True)
        m = re.search(r"Název materiálu\s*\n(.+)", txt)
        if m:
            nazev = m.group(1).strip()
        m = re.search(r"Čj\.\s*\n(.+)", txt)
        if m:
            cj = m.group(1).strip()
        m = re.search(r"PID\s*\n(\S+)", txt)
        if m:
            pid = m.group(1).strip()
    base = ODOK + "/portal/services/download/attachment/" + str(year) + "/" + str(cislo) + "/"
    text, src = "", ""
    data = get(base + "docx/", binary=True)
    if data and data[:2] == b"PK":
        try:
            text = clean(docx_text(data))
            src = base + "docx/"
        except Exception as e:  # noqa
            log("  ! docx " + str(cislo) + ": " + str(e))
    if len(text) < 40:
        data = get(base + "pdf/", binary=True)
        if data and data[:4] == b"%PDF":
            try:
                text = clean(pdf_text(data))
                src = base + "pdf/"
            except Exception as e:  # noqa
                log("  ! pdf " + str(cislo) + ": " + str(e))
    return {
        "organ": "vlada", "cislo": cislo, "rok": year, "datum": datum, "volebni_obdobi": "",
        "schuze": None, "nazev": nazev or None, "material_cj": cj or None, "material_pid": pid or None,
        "zdroj_url": ODOK + "/portal/zvlady/usneseni/" + str(year) + "/" + str(cislo) + "/",
        "soubor_url": src or None, "text_plny": text,
    }


# ---------------------------------------------------------------- PS
def ps_parse_date(s):
    m = re.search(r"\((\d{1,2})\.\s*([a-záčďéěíňóřšťúůýž]+)\s+(\d{4})\)\s*(?:\(přílohy\))?\s*$", s.strip())
    if not m or m.group(2) not in MONTHS:
        return None
    return dt.date(int(m.group(3)), MONTHS[m.group(2)], int(m.group(1)))


def ps_list(term):
    """seznam usneseni PS pro volebni obdobi (cislo, nazev, datum, idd)"""
    out = []
    page = 1
    last = 1
    while page <= last and time_left() > 600:
        url = PSP + "/sqw/hp.sqw?k=99&o=" + str(term) + "&td=8" + ("&n=" + str(page) if page > 1 else "")
        html = get(url)
        if not html:
            break
        soup = BeautifulSoup(html, "html.parser")
        for a in soup.find_all("a", href=True):
            m = re.search(r"[?&]n=(\d+)", a["href"])
            if m and "td=8" in a["href"]:
                last = max(last, int(m.group(1)))
        for a in soup.find_all("a", href=True):
            m = re.search(r"text2\.sqw\?idd=(\d+)", a["href"])
            mc = re.match(r"č\.\s*(\d+)", a.get_text(strip=True))
            if not m or not mc:
                continue
            row = a.parent
            raw = row.get_text(" ", strip=True) if row else a.get_text()
            raw = re.sub(r"^č\.\s*\d+\s*", "", raw)
            out.append((int(mc.group(1)), raw, ps_parse_date(raw), int(m.group(1))))
        page += 1
    log("PS obdobi " + str(term) + ": nalezeno " + str(len(out)) + " usneseni, stran " + str(last))
    return out


def ps_fetch(term, cislo, nazev_raw, datum, idd):
    nazev = re.sub(r"\s*\(\d{1,2}\.\s*[a-záčďéěíňóřšťúůýž]+\s+\d{4}\)\s*(\(přílohy\))?\s*$", "", nazev_raw).strip()
    url = PSP + "/sqw/text/orig2.sqw?idd=" + str(idd)
    text, src = "", ""
    data = get(url, binary=True)
    if data and data[:2] == b"PK":
        try:
            text = clean(docx_text(data))
            src = url
        except Exception as e:  # noqa
            log("  ! docx PS " + str(cislo) + ": " + str(e))
    if len(text) < 40:
        data = get(PSP + "/sqw/text/orig2.sqw?idd=" + str(idd + 1), binary=True)
        if data and data[:4] == b"%PDF":
            try:
                text = clean(pdf_text(data))
                src = PSP + "/sqw/text/orig2.sqw?idd=" + str(idd + 1)
            except Exception as e:  # noqa
                log("  ! pdf PS " + str(cislo) + ": " + str(e))
    return {
        "organ": "ps", "cislo": cislo, "rok": datum.year, "datum": datum, "volebni_obdobi": str(term),
        "schuze": None, "nazev": nazev or None, "material_cj": None, "material_pid": None,
        "zdroj_url": PSP + "/sqw/text/text2.sqw?idd=" + str(idd), "soubor_url": src or None,
        "text_plny": text,
    }


# ---------------------------------------------------------------- DB
def save(conn, doc):
    h = hashlib.sha256((doc["text_plny"] + "|" + (doc["nazev"] or "")).encode("utf-8")).hexdigest()
    with conn.cursor() as cur:
        cur.execute(
            "select id, content_hash from documents where organ=%s and rok=%s and cislo=%s and volebni_obdobi=%s",
            (doc["organ"], doc["rok"], doc["cislo"], doc["volebni_obdobi"]),
        )
        row = cur.fetchone()
        if row and row[1] == h:
            return "beze_zmeny"
        title = ("Usnesení vlády č. " if doc["organ"] == "vlada" else "Usnesení PS č. ") + str(doc["cislo"]) + "/" + str(doc["rok"])
        if doc["nazev"]:
            title += " – " + doc["nazev"]
        chunks = make_chunks(title, doc["text_plny"])
        if row:
            doc_id = row[0]
            cur.execute(
                "update documents set datum=%s, schuze=%s, nazev=%s, material_cj=%s, material_pid=%s, zdroj_url=%s, "
                "soubor_url=%s, text_plny=%s, content_hash=%s, updated_at=now() where id=%s",
                (doc["datum"], doc["schuze"], doc["nazev"], doc["material_cj"], doc["material_pid"],
                 doc["zdroj_url"], doc["soubor_url"], doc["text_plny"], h, doc_id),
            )
            cur.execute("delete from chunks where document_id=%s", (doc_id,))
            result = "aktualizovano"
        else:
            cur.execute(
                "insert into documents (organ,cislo,rok,datum,volebni_obdobi,schuze,nazev,material_cj,material_pid,"
                "zdroj_url,soubor_url,text_plny,content_hash) values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) returning id",
                (doc["organ"], doc["cislo"], doc["rok"], doc["datum"], doc["volebni_obdobi"], doc["schuze"],
                 doc["nazev"], doc["material_cj"], doc["material_pid"], doc["zdroj_url"], doc["soubor_url"],
                 doc["text_plny"], h),
            )
            doc_id = cur.fetchone()[0]
            result = "nove"
        for i, c in enumerate(chunks):
            cur.execute("insert into chunks (document_id, chunk_index, content) values (%s,%s,%s)", (doc_id, i, c))
    conn.commit()
    return result


def existing_keys(conn):
    with conn.cursor() as cur:
        cur.execute("select organ, rok, cislo, volebni_obdobi, content_hash is not null and length(text_plny) >= 40 from documents")
        return {(a, b, c, d): e for a, b, c, d, e in cur.fetchall()}


def main():
    url = os.environ.get("NEON_ROZHODNICEK_DB_URL")
    if not url and not DRY:
        log("Chybi NEON_ROZHODNICEK_DB_URL")
        sys.exit(1)
    years = [int(x) for x in (os.environ.get("ROZHODNICEK_YEARS") or str(dt.date.today().year)).replace(" ", "").split(",") if x]
    sources = [x.strip() for x in (os.environ.get("ROZHODNICEK_SOURCES") or "vlada,ps").split(",")]
    conn = psycopg2.connect(url) if url else None
    have = existing_keys(conn) if conn else {}
    stats = {"nove": 0, "aktualizovano": 0, "beze_zmeny": 0, "bez_textu": 0, "chyby": 0}
    log("Roky: " + str(years) + ", zdroje: " + str(sources) + (" (DRY)" if DRY else ""))

    def handle(doc):
        if not doc["text_plny"] or len(doc["text_plny"]) < 40:
            stats["bez_textu"] += 1
            log("  ! bez textu: " + doc["organ"] + " " + str(doc["cislo"]) + "/" + str(doc["rok"]))
            if DRY or not doc["nazev"]:
                return
        if DRY:
            stats["nove"] += 1
            if stats["nove"] <= 3:
                log("  ukazka: " + doc["organ"] + " " + str(doc["cislo"]) + "/" + str(doc["rok"]) + " | " + str(doc["nazev"])[:100]
                    + " | text " + str(len(doc["text_plny"])) + " zn.: " + doc["text_plny"][:160].replace("\n", " / "))
            return
        try:
            stats[save(conn, doc)] += 1
        except Exception as e:  # noqa
            stats["chyby"] += 1
            conn.rollback()
            log("  ! DB chyba " + doc["organ"] + " " + str(doc["cislo"]) + ": " + str(e))

    if "vlada" in sources:
        for year in years:
            log("== Vlada " + str(year))
            for cislo, datum in vlada_resolutions(year):
                if time_left() < 300:
                    log("Dosazen casovy limit"); break
                if have.get(("vlada", year, cislo, "")):
                    stats["beze_zmeny"] += 1
                    continue
                handle(vlada_fetch(year, cislo, datum))
    if "ps" in sources:
        terms = [9, 10] if min(years) <= 2025 else [10]
        for term in terms:
            log("== PS volebni obdobi " + str(term))
            for cislo, nazev_raw, datum, idd in ps_list(term):
                if datum is None or datum.year not in years:
                    continue
                if time_left() < 300:
                    log("Dosazen casovy limit"); break
                if have.get(("ps", datum.year, cislo, str(term))):
                    stats["beze_zmeny"] += 1
                    continue
                handle(ps_fetch(term, cislo, nazev_raw, datum, idd))
    log("=== Hotovo: " + str(stats) + " ===")
    if conn:
        conn.close()


if __name__ == "__main__":
    main()
